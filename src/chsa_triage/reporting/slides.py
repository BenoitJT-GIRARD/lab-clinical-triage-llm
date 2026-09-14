"""Support de soutenance, généré à partir des résultats.

La présentation dure quinze minutes et sert à démontrer la faisabilité technique
et la valeur clinique du prototype. Elle suit donc la trame d'une démonstration,
pas celle du rapport : le problème, ce qu'on a trouvé en ouvrant les données, ce
qu'on a construit, ce que ça donne, et ce qu'il faudrait pour aller plus loin.

Comme le rapport, elle est produite depuis les fichiers de résultats : les
chiffres de la soutenance sont ceux du dépôt, à la décimale près.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

BLEU = RGBColor(0x1F, 0x4E, 0x79)
ENCRE = RGBColor(0x1A, 0x1A, 0x19)
ENCRE_SECONDAIRE = RGBColor(0x52, 0x51, 0x4E)

LARGEUR = Cm(33.87)
HAUTEUR = Cm(19.05)


@dataclass
class Slide:
    """Une diapositive : un titre, des puces, et éventuellement une figure."""

    titre: str
    puces: list[str] = field(default_factory=list)
    figure: Path | None = None
    note: str = ""


def _ajouter_zone_texte(diapositive, gauche, haut, largeur, hauteur):
    zone = diapositive.shapes.add_textbox(gauche, haut, largeur, hauteur)
    cadre = zone.text_frame
    cadre.word_wrap = True
    return cadre


def _page_de_titre(presentation: Presentation, titre: str, sous_titre: str, auteur: str) -> None:
    diapositive = presentation.slides.add_slide(presentation.slide_layouts[6])
    cadre = _ajouter_zone_texte(diapositive, Cm(3), Cm(6), LARGEUR - Cm(6), Cm(7))

    paragraphe = cadre.paragraphs[0]
    paragraphe.text = titre
    paragraphe.alignment = PP_ALIGN.CENTER
    paragraphe.runs[0].font.size = Pt(40)
    paragraphe.runs[0].font.bold = True
    paragraphe.runs[0].font.color.rgb = BLEU

    for texte, taille in ((sous_titre, 20), (auteur, 16)):
        suivant = cadre.add_paragraph()
        suivant.text = texte
        suivant.alignment = PP_ALIGN.CENTER
        suivant.space_before = Pt(14)
        suivant.runs[0].font.size = Pt(taille)
        suivant.runs[0].font.color.rgb = ENCRE_SECONDAIRE


def _diapositive(presentation: Presentation, contenu: Slide) -> None:
    diapositive = presentation.slides.add_slide(presentation.slide_layouts[6])

    cadre_titre = _ajouter_zone_texte(diapositive, Cm(1.6), Cm(1.0), LARGEUR - Cm(3.2), Cm(2))
    paragraphe = cadre_titre.paragraphs[0]
    paragraphe.text = contenu.titre
    paragraphe.runs[0].font.size = Pt(28)
    paragraphe.runs[0].font.bold = True
    paragraphe.runs[0].font.color.rgb = BLEU

    largeur_texte = LARGEUR - Cm(3.2) if contenu.figure is None else Cm(14)
    if contenu.puces:
        cadre = _ajouter_zone_texte(diapositive, Cm(1.6), Cm(3.6), largeur_texte, Cm(12))
        for index, puce in enumerate(contenu.puces):
            paragraphe = cadre.paragraphs[0] if index == 0 else cadre.add_paragraph()
            paragraphe.text = f"•  {puce}"
            paragraphe.space_after = Pt(12)
            paragraphe.runs[0].font.size = Pt(17)
            paragraphe.runs[0].font.color.rgb = ENCRE

    if contenu.figure is not None and contenu.figure.exists():
        gauche = Cm(16.5) if contenu.puces else Cm(3)
        largeur = Cm(15.5) if contenu.puces else LARGEUR - Cm(6)
        diapositive.shapes.add_picture(str(contenu.figure), gauche, Cm(4.2), width=largeur)

    if contenu.note:
        diapositive.notes_slide.notes_text_frame.text = contenu.note


def _proprietes(presentation, titre: str, sous_titre: str, auteur: str) -> None:
    """Renseigne les propriétés du fichier, que le lecteur affiche et que l'on cite.

    Sans cela, le document part avec celles du gabarit vide de la bibliothèque de
    génération : le nom de son auteur, une mention d'outil et des dates de 2013.
    """
    from datetime import UTC, datetime

    maintenant = datetime.now(UTC).replace(tzinfo=None)
    # La ligne de titre porte le nom, le rôle et la date ; les propriétés du
    # fichier n'attendent que le nom.
    nom = auteur.split("·")[0].strip()
    proprietes = presentation.core_properties
    proprietes.author = nom
    proprietes.last_modified_by = nom
    proprietes.title = titre
    proprietes.subject = sous_titre
    proprietes.comments = ""
    proprietes.category = ""
    proprietes.created = maintenant
    proprietes.modified = maintenant


def construire(
    chemin: Path,
    titre: str,
    sous_titre: str,
    auteur: str,
    diapositives: list[Slide],
) -> Path:
    """Assemble la présentation et l'enregistre."""
    presentation = Presentation()
    presentation.slide_width = LARGEUR
    presentation.slide_height = HAUTEUR

    _proprietes(presentation, titre, sous_titre, auteur)
    _page_de_titre(presentation, titre, sous_titre, auteur)
    for contenu in diapositives:
        _diapositive(presentation, contenu)

    chemin.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(chemin))
    return chemin
